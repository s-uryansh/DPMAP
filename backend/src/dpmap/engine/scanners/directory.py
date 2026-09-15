"""Bounded, read-only scanning for service-visible directories."""

import csv
from dataclasses import dataclass
import os
from pathlib import Path
from typing import Iterable, Iterator
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from dpmap.engine.detectors import aggregate_chunks


SUPPORTED_SUFFIXES = frozenset({".txt", ".csv", ".xlsx"})
TEXT_CHUNK_CHARS = 64 * 1024
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_CSV_FIELD_CHARS = 1024 * 1024
MAX_XLSX_ENTRIES = 10_000
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
csv.field_size_limit(MAX_CSV_FIELD_CHARS)


class DirectoryScanError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class InventoryAggregate:
    field_locator: str | None
    pii_type: str
    match_count: int
    units_with_pii: int
    units_scanned: int
    confidence_band: str
    reason_codes: tuple[str, ...]


@dataclass(frozen=True)
class FileScanResult:
    path: Path
    items: tuple[InventoryAggregate, ...] = ()
    units_scanned: int = 0
    bytes_scanned: int = 0
    warnings: tuple[str, ...] = ()
    scanned: bool = False


def configured_roots(value: str | None) -> tuple[Path, ...]:
    if not value:
        raise DirectoryScanError(
            "scan_roots_not_configured", "Allowed scan roots are not configured"
        )
    try:
        roots = tuple(
            Path(item).expanduser().resolve(strict=True)
            for item in value.split(os.pathsep)
            if item
        )
    except OSError as error:
        raise DirectoryScanError(
            "invalid_scan_root", "An allowed scan root is unavailable"
        ) from error
    if not roots or any(not root.is_dir() for root in roots):
        raise DirectoryScanError(
            "invalid_scan_root", "Allowed scan roots must be directories"
        )
    return roots


def validate_directory(path: str, allowed_roots: Iterable[Path]) -> Path:
    requested = Path(path)
    if not requested.is_absolute():
        raise DirectoryScanError("invalid_path", "Directory path must be absolute")
    lexical = Path(os.path.abspath(requested))
    root = next(
        (candidate for candidate in allowed_roots if lexical.is_relative_to(candidate)),
        None,
    )
    if root is None:
        raise DirectoryScanError(
            "path_not_allowed", "Directory is outside the allowed scan roots"
        )

    current = root
    for part in lexical.relative_to(root).parts:
        current /= part
        if current.is_symlink():
            raise DirectoryScanError(
                "symlink_target", "Directory targets cannot contain symlinks"
            )
    try:
        resolved = lexical.resolve(strict=True)
    except FileNotFoundError as error:
        raise DirectoryScanError(
            "path_not_found", "Directory does not exist"
        ) from error
    except OSError as error:
        raise DirectoryScanError(
            "path_unavailable", "Directory cannot be accessed"
        ) from error
    if not resolved.is_dir():
        raise DirectoryScanError("not_a_directory", "Target is not a directory")
    if not os.access(resolved, os.R_OK | os.X_OK):
        raise DirectoryScanError(
            "permission_denied", "Directory is not readable"
        )
    return resolved


def _confidence_band(confidence: float) -> str:
    if confidence >= 0.95:
        return "high"
    if confidence >= 0.8:
        return "medium"
    return "low"


def _items_from_summary(
    summary: dict[str, object],
    *,
    field_locator: str | None,
    units_scanned: int,
) -> list[InventoryAggregate]:
    items = []
    for pii_type, match in summary["matches"].items():
        items.append(
            InventoryAggregate(
                field_locator=field_locator,
                pii_type=pii_type,
                match_count=match["match_count"],
                units_with_pii=match["units_with_pii"],
                units_scanned=units_scanned,
                confidence_band=_confidence_band(match["confidence"]),
                reason_codes=tuple(match["reason_codes"]),
            )
        )
    return items


def _text_chunks(path: Path, encoding: str) -> Iterator[str]:
    with path.open("r", encoding=encoding, errors="strict") as source:
        while chunk := source.read(TEXT_CHUNK_CHARS):
            yield chunk


def _scan_txt(path: Path, detectors: frozenset[str]) -> tuple[list, int, list]:
    warnings = []
    try:
        summary = aggregate_chunks(
            _text_chunks(path, "utf-8-sig"), enabled=detectors
        )
    except UnicodeDecodeError:
        summary = aggregate_chunks(_text_chunks(path, "cp1252"), enabled=detectors)
        warnings.append("encoding_fallback_cp1252")
    units = summary["chunks_scanned"]
    return _items_from_summary(
        summary, field_locator=None, units_scanned=units
    ), units, warnings


def _csv_reader(path: Path, encoding: str):
    source = path.open("r", encoding=encoding, errors="strict", newline="")
    try:
        sample = source.read(8192)
        source.seek(0)
        if not sample:
            return source, csv.reader(source, csv.excel), False
        sniffer = csv.Sniffer()
        try:
            dialect = sniffer.sniff(sample)
        except csv.Error:
            dialect = csv.excel
        try:
            has_header = sniffer.has_header(sample)
        except csv.Error:
            has_header = False
        return source, csv.reader(source, dialect), has_header
    except Exception:
        source.close()
        raise


def _scan_csv_encoding(
    path: Path, detectors: frozenset[str], encoding: str
) -> tuple[list[InventoryAggregate], int]:
    source, rows, has_header = _csv_reader(path, encoding)
    aggregates: dict[tuple[str, str], dict[str, object]] = {}
    units_by_field: dict[str, int] = {}
    try:
        headers = next(rows, ()) if has_header else ()
        row_count = 0
        for row in rows:
            row_count += 1
            for index, value in enumerate(row):
                field = f"column:{index + 1}"
                units_by_field[field] = units_by_field.get(field, 0) + 1
                context = headers[index] if index < len(headers) else ""
                summary = aggregate_chunks(
                    [value], enabled=detectors, context=context
                )
                for pii_type, match in summary["matches"].items():
                    key = (field, pii_type)
                    item = aggregates.setdefault(
                        key,
                        {
                            "matches": 0,
                            "units": 0,
                            "confidence": 0.0,
                            "reasons": set(),
                        },
                    )
                    item["matches"] += match["match_count"]
                    item["units"] += 1
                    item["confidence"] = max(
                        item["confidence"], match["confidence"]
                    )
                    item["reasons"].update(match["reason_codes"])
    finally:
        source.close()

    items = [
        InventoryAggregate(
            field_locator=field,
            pii_type=pii_type,
            match_count=item["matches"],
            units_with_pii=item["units"],
            units_scanned=units_by_field[field],
            confidence_band=_confidence_band(item["confidence"]),
            reason_codes=tuple(sorted(item["reasons"])),
        )
        for (field, pii_type), item in sorted(aggregates.items())
    ]
    return items, row_count


def _scan_csv(path: Path, detectors: frozenset[str]) -> tuple[list, int, list]:
    try:
        items, units = _scan_csv_encoding(path, detectors, "utf-8-sig")
        return items, units, []
    except UnicodeDecodeError:
        items, units = _scan_csv_encoding(path, detectors, "cp1252")
        return items, units, ["encoding_fallback_cp1252"]


def _check_xlsx_archive(path: Path) -> None:
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_XLSX_ENTRIES:
            raise DirectoryScanError(
                "xlsx_too_many_entries", "XLSX has too many entries"
            )
        if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise DirectoryScanError("xlsx_too_large", "XLSX expands beyond the limit")


def _scan_xlsx(path: Path, detectors: frozenset[str]) -> tuple[list, int, list]:
    _check_xlsx_archive(path)
    workbook = load_workbook(
        path, read_only=True, data_only=True, keep_links=False
    )
    aggregates: dict[tuple[str, str], dict[str, object]] = {}
    units_by_field: dict[str, int] = {}
    units = 0
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
            for row in worksheet.iter_rows(values_only=True):
                for column_index, value in enumerate(row, 1):
                    if value is None:
                        continue
                    text = str(value)
                    if len(text) > MAX_CSV_FIELD_CHARS:
                        raise DirectoryScanError(
                            "xlsx_cell_too_large", "XLSX cell exceeds the limit"
                        )
                    units += 1
                    field = f"sheet:{sheet_index}/column:{column_index}"
                    units_by_field[field] = units_by_field.get(field, 0) + 1
                    summary = aggregate_chunks([text], enabled=detectors)
                    for pii_type, match in summary["matches"].items():
                        key = (field, pii_type)
                        item = aggregates.setdefault(
                            key,
                            {
                                "matches": 0,
                                "units": 0,
                                "confidence": 0.0,
                                "reasons": set(),
                            },
                        )
                        item["matches"] += match["match_count"]
                        item["units"] += 1
                        item["confidence"] = max(
                            item["confidence"], match["confidence"]
                        )
                        item["reasons"].update(match["reason_codes"])
    finally:
        workbook.close()

    items = [
        InventoryAggregate(
            field_locator=field,
            pii_type=pii_type,
            match_count=item["matches"],
            units_with_pii=item["units"],
            units_scanned=units_by_field[field],
            confidence_band=_confidence_band(item["confidence"]),
            reason_codes=tuple(sorted(item["reasons"])),
        )
        for (field, pii_type), item in sorted(aggregates.items())
    ]
    return items, units, []


def _scan_file(path: Path, detectors: frozenset[str]) -> FileScanResult:
    try:
        before = path.stat()
        if before.st_size > MAX_FILE_BYTES:
            return FileScanResult(path, warnings=("file_too_large",))
        scanner = {
            ".txt": _scan_txt,
            ".csv": _scan_csv,
            ".xlsx": _scan_xlsx,
        }[path.suffix.lower()]
        items, units, warnings = scanner(path, detectors)
        after = path.stat()
        if (before.st_size, before.st_mtime_ns) != (
            after.st_size,
            after.st_mtime_ns,
        ):
            warnings.append("unstable_file")
        return FileScanResult(
            path,
            tuple(items),
            units,
            before.st_size,
            tuple(warnings),
            True,
        )
    except DirectoryScanError as error:
        return FileScanResult(path, warnings=(error.code,))
    except (
        BadZipFile,
        csv.Error,
        InvalidFileException,
        OSError,
        UnicodeError,
        ValueError,
    ):
        return FileScanResult(path, warnings=("file_unreadable",))


def scan_directory(
    root: Path, detectors: Iterable[str]
) -> Iterator[FileScanResult]:
    selected = frozenset(detectors)
    directories = [root]
    while directories:
        directory = directories.pop()
        try:
            entries = sorted(os.scandir(directory), key=lambda entry: entry.name)
        except OSError:
            yield FileScanResult(directory, warnings=("directory_unreadable",))
            continue

        children = []
        for entry in entries:
            path = Path(entry.path)
            try:
                if entry.is_symlink():
                    yield FileScanResult(path, warnings=("symlink_skipped",))
                elif entry.name.startswith("."):
                    yield FileScanResult(path, warnings=("hidden_path_skipped",))
                elif entry.is_dir(follow_symlinks=False):
                    children.append(path)
                elif entry.is_file(follow_symlinks=False):
                    if path.suffix.lower() in SUPPORTED_SUFFIXES:
                        yield _scan_file(path, selected)
                    else:
                        yield FileScanResult(
                            path, warnings=("unsupported_file_skipped",)
                        )
            except OSError:
                yield FileScanResult(path, warnings=("path_unreadable",))
        directories.extend(reversed(children))
